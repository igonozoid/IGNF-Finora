from datetime import date
from decimal import Decimal as D

from openpyxl import load_workbook

from finora.services import export


def _sheet():
    s = export.Sheet("DRE pessoal", "Últimos 3 meses", ["Linha", "Set", "Total", "Qtd."])
    s.add(["Receitas", D("5000.00"), D("5000.00"), 2], "bold")
    s.add(["Aluguel", D("-1500.50"), D("-1500.50"), 1], "detail")
    s.add(["Vence em", date(2026, 10, 6), None, 0])
    s.footer = "Rodapé"
    return s


def test_excel_grava_numeros_de_verdade(tmp_path):
    path = tmp_path / "x.xlsx"
    export.to_xlsx(_sheet(), str(path))
    ws = load_workbook(path).active
    assert ws["A1"].value == "DRE pessoal" and ws["A2"].value == "Últimos 3 meses"
    assert [c.value for c in ws[4]] == ["Linha", "Set", "Total", "Qtd."]
    assert ws["B5"].value == 5000 and ws["B6"].value == -1500.5 and "R$" in ws["B6"].number_format
    assert ws["A5"].font.bold and ws["D5"].value == 2
    assert ws["B7"].is_date and ws["A9"].value == "Rodapé"


def test_html_para_pdf():
    html = export.to_html(_sheet())
    assert "DRE pessoal" in html and "R$ 5.000,00" in html and "− R$ 1.500,50" in html
    assert "06/10/2026" in html and 'align="right"' in html and "Rodapé" in html


def test_nome_de_arquivo_seguro():
    assert export.safe_name("Extrato por conta — Nubank · Este mês") == "Extrato por conta - Nubank - Este mês"
    assert export.safe_name("a/b:c?") == "a-b-c"
    assert export.safe_name("///") == "Finora"
