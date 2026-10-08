from openpyxl import load_workbook

from tests.ui.conftest import pump


def test_exportar_lancamentos_e_relatorios(plus, window, dialogs, tmp_path, qtbot):
    window.go_to("entries")
    page = window.page("entries")
    assert page.export_btn.menu() is not None                    # Plus: menu Excel / PDF
    dialogs.save_path = str(tmp_path / "lanc")                   # sem extensão: o app completa
    page._export("xlsx")
    ws = load_workbook(tmp_path / "lanc.xlsx").active
    descr = [ws.cell(r, 3).value for r in range(5, ws.max_row + 1)]
    assert "Salário" in descr and "Mercado" in descr

    window.go_to("reports")
    exported = 0
    for key, view in window.page("reports").views.items():
        view.refresh()
        pump(qtbot)
        sheet = view.export_sheet()
        if sheet is None or not sheet.rows:
            continue
        dialogs.save_path = str(tmp_path / f"{key}.pdf")
        assert [b.property("fullText") for b in view.export_btns] == ["Imprimir", "Excel", "PDF"], key
        view.export_btns[2].click()                              # botão "PDF"
        assert (tmp_path / f"{key}.pdf").read_bytes()[:4] == b"%PDF", key
        exported += 1
    assert exported >= 5


def test_exportar_bloqueado_na_free(window, dialogs):
    page = window.page("entries")
    actions = page.export_btn.menu().actions()                   # Imprimir liberado; Excel e PDF com convite
    assert actions[0].text() == "Imprimir…" and "Plus" in actions[2].text()
    actions[2].trigger()
    assert dialogs.shown, "deveria convidar para o upgrade"


def test_imprimir_em_todos_os_relatorios_na_free(window, qtbot, monkeypatch):
    """Todo relatório tem Imprimir, também na Free, com a mesma janela: prévia, retrato/paisagem e cabeçalho."""
    from finora.core import settings
    from finora.ui import printing
    from finora.ui.reports_page import REPORTS
    seen = []

    def fake_exec(dlg):
        dlg.preview.updatePreview()
        assert dlg.preview.pageCount() >= 1
        land = dlg.landscape
        dlg.set_landscape(not land)                              # troca e lembra para a próxima vez
        assert settings.get_print_landscape(dlg.sheet.title) is (not land)
        assert not dlg.can_pdf and dlg.sheet.head.name == window.profile.name
        seen.append(dlg.sheet.title)
        return 0

    monkeypatch.setattr(printing.PrintDialog, "exec", fake_exec)
    window.go_to("reports")
    views = window.page("reports").views
    for key, _label, _icon, feature in REPORTS:
        if feature:                                              # relatório pago: nem abre na Free
            continue
        view = views[key]
        view.refresh()
        assert view.print_btn.isEnabled(), key
        if view.export_sheet() is not None and view.export_sheet().rows:
            view.print_btn.click()
    assert len(seen) >= 5
    window.go_to("entries")
    window.page("entries").export_btn.menu().actions()[0].trigger()
    assert seen[-1] == "Lançamentos"


def test_cabecalho_da_entidade_igual_ao_recibo(plus, window, dialogs, tmp_path):
    from finora.core import db
    from finora.services import entity_admin
    with db.Session() as s:
        entity_admin.update_details(s, window.profile.id, city="Teutônia", state="RS", phone="(51) 98101-6345")
    window.go_to("entries")
    page = window.page("entries")
    dialogs.save_path = str(tmp_path / "lanc.xlsx")
    page._export("xlsx")
    ws = load_workbook(tmp_path / "lanc.xlsx").active
    assert ws["A1"].value == window.profile.name and "Teutônia/RS" in ws["A2"].value
    from finora.services import export
    from finora.ui import exporting
    sheet = exporting.prepare(page.export_sheet())
    html = export.to_html(sheet)
    assert "Teutônia/RS" in html and "(51) 98101-6345" in html


def test_a_pagar_e_receber_periodo_personalizado(window, qtbot):
    from datetime import date
    window.go_to("reports")
    view = window.page("reports").views["agenda"]
    view.period.setCurrentIndex(view.period.findData("custom"))
    assert not view.first.isHidden() and not view.last.isHidden()
    view.first.setDate(date(2000, 1, 1))
    view.last.setDate(date(2100, 12, 31))
    view.kind.setCurrentIndex(view.kind.findData("expense"))      # só a pagar: some a coluna A RECEBER
    assert [c[0] for c in view.COLUMNS][-1] == "A PAGAR" and "A RECEBER" not in [c[0] for c in view.COLUMNS]
    sheet = view.export_sheet()
    assert "01/01/2000 a 31/12/2100" in sheet.subtitle and "Só a pagar" in sheet.subtitle
    assert "A receber" not in view.footer.text()
    view.first.setDate(date(2101, 1, 1))
    assert "inicial" in view.empty.text()
