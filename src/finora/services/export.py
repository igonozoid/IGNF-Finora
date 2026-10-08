"""Exportar tabelas (lançamentos e relatórios) para Excel (.xlsx) e para HTML, que a tela imprime em PDF."""
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from html import escape

from finora.core import money

# estilos de linha: "" normal · "bold" subtotal/título · "detail" linha de detalhe (recuada, cinza) · "result" resultado
STYLES = ("", "bold", "detail", "result")
_PERCENT = re.compile(r"^-?[\d.,]+%$")


@dataclass
class Sheet:
    title: str
    subtitle: str
    headers: list[str]
    rows: list[list] = field(default_factory=list)      # célula: str, int, Decimal, date ou None
    styles: list[str] = field(default_factory=list)     # um por linha (vazio = normal)
    currency: str = "BRL"
    footer: str = ""
    head: object = None          # entity_admin.Letterhead: cabeçalho com nome, endereço e logotipo da entidade

    def add(self, cells: list, style: str = "") -> None:
        assert len(cells) == len(self.headers) and style in STYLES
        self.rows.append(list(cells))
        self.styles.append(style)

    def style(self, i: int) -> str:
        return self.styles[i] if i < len(self.styles) else ""

    def numeric_cols(self) -> set[int]:
        """Colunas que têm números (alinhadas à direita)."""
        def num(v) -> bool:
            if isinstance(v, str):
                return bool(_PERCENT.match(v))
            return isinstance(v, (int, Decimal)) and not isinstance(v, bool)
        return {c for c in range(len(self.headers)) if any(num(r[c]) for r in self.rows)}


def safe_name(text: str) -> str:
    """Sugestão de nome de arquivo: 'DRE pessoal — 2026' -> 'DRE pessoal - 2026'."""
    out = "".join(ch if ch.isalnum() or ch in " -_.," else "-" for ch in text.replace("—", "-"))
    return " ".join(out.split()).strip(" .-") or "Finora"


# ---------- Excel ----------
def to_xlsx(sheet: Sheet, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = safe_name(sheet.title)[:31]
    if sheet.head is not None:                   # nome da entidade na primeira linha
        ws.append([sheet.head.name])
        ws["A1"].font = Font(bold=True, size=11, color="444444")
    ws.append([sheet.title])
    first = ws.max_row
    ws.cell(first, 1).font = Font(bold=True, size=13)
    ws.append([sheet.subtitle])
    ws.cell(first + 1, 1).font = Font(italic=True, color="666666")
    ws.append([])
    ws.append(sheet.headers)
    head = ws.max_row
    for c in range(1, len(sheet.headers) + 1):
        cell = ws.cell(head, c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEEEEE")
    sym = money.symbol(sheet.currency).replace('"', "")
    money_fmt = f'"{sym}" #,##0.00;[Red]-"{sym}" #,##0.00'
    numeric = sheet.numeric_cols()
    for i, row in enumerate(sheet.rows):
        ws.append(row)        # Decimal vai como número (o openpyxl não passa por float)
        r = ws.max_row
        style = sheet.style(i)
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c)
            if isinstance(v, Decimal):
                cell.number_format = money_fmt
            elif isinstance(v, (date, datetime)):
                cell.number_format = "DD/MM/YYYY"
            if c - 1 in numeric:
                cell.alignment = Alignment(horizontal="right")
            if style in ("bold", "result"):
                cell.font = Font(bold=True)
            elif style == "detail":
                cell.font = Font(color="777777")
                if c == 1:
                    cell.alignment = Alignment(indent=2)
    if sheet.footer:
        ws.append([])
        ws.append([sheet.footer])
    for c in range(1, len(sheet.headers) + 1):
        longest = max([len(str(sheet.headers[c - 1]))] + [len(_text(r[c - 1], sheet.currency)) for r in sheet.rows])
        ws.column_dimensions[get_column_letter(c)].width = min(60, max(10, longest + 2))
    ws.freeze_panes = ws.cell(head + 1, 1)
    wb.save(path)


# ---------- HTML (para PDF) ----------
def _text(v, currency: str) -> str:
    if v is None:
        return ""
    if isinstance(v, Decimal):
        return money.fmt(v, currency)
    if isinstance(v, (date, datetime)):
        return v.strftime("%d/%m/%Y")
    return str(v)


LOGO_URL = "finora-logo"


def _letterhead_html(head) -> str:
    """Cabeçalho da entidade (o mesmo do recibo). O logotipo é um recurso do documento (ui/exporting)."""
    if head is None:
        return ""
    lines = "".join(f'<br><span style="font-size:7pt; color:#444">{escape(x)}</span>'
                    for x in (head.address_line, head.contacts_line, head.docs_line) if x)
    logo = f'<td width="70" valign="middle"><img src="{LOGO_URL}" height="36"></td>' if head.logo else ""
    return ('<table width="100%" cellspacing="0" cellpadding="0" style="margin-bottom:6px"><tr>'
            f'{logo}<td valign="middle"><b style="font-size:10pt">{escape(head.name)}</b>{lines}</td></tr></table>'
            '<table width="100%" cellspacing="0" cellpadding="0" style="margin:0 0 8px 0"><tr>'
            '<td style="border-top:1px solid #999; font-size:2pt">&nbsp;</td></tr></table>')


def to_html(sheet: Sheet, generated: datetime | None = None) -> str:
    generated = generated or datetime.now()
    numeric = sheet.numeric_cols()
    # a coluna de texto mais "longa" (descrição, nome) fica com a sobra de largura
    texts = [c for c in range(len(sheet.headers)) if c not in numeric]
    wide = max(texts, key=lambda c: sum(len(_text(r[c], sheet.currency)) for r in sheet.rows), default=None)
    head = "".join(f'<th align="{"right" if c in numeric else "left"}"{" width=\"34%\"" if c == wide else ""}>'
                   f'{escape(h)}</th>' for c, h in enumerate(sheet.headers))
    body = []
    for i, row in enumerate(sheet.rows):
        style = sheet.style(i)
        tr = {"bold": ' style="font-weight:bold; background:#f0f0f0"',
              "result": ' style="font-weight:bold; background:#e8e8e8"',
              "detail": ' style="color:#666"'}.get(style, "")
        cells = []
        for c, v in enumerate(row):
            text = "–" if isinstance(v, Decimal) and v == 0 else escape(_text(v, sheet.currency))
            if style == "detail" and c == 0:
                text = "&nbsp;&nbsp;&nbsp;&nbsp;" + text
            neg = isinstance(v, Decimal) and v < 0
            color = ' style="color:#b03020"' if neg else ""
            cells.append(f'<td align="{"right" if c in numeric else "left"}"{color}>{text}</td>')
        body.append(f"<tr{tr}>{''.join(cells)}</tr>")
    foot = f'<p style="font-size:8pt">{escape(sheet.footer)}</p>' if sheet.footer else ""
    return (
        "<html><body style=\"font-family:'IBM Plex Sans',Arial,Helvetica,sans-serif; font-size:8pt; color:#000\">"
        f'{_letterhead_html(sheet.head)}<h2 style="margin:0">{escape(sheet.title)}</h2>'
        f'<p style="color:#555; margin:2px 0 8px 0">{escape(sheet.subtitle)}</p>'
        '<table width="100%" cellspacing="0" cellpadding="3" border="0" style="border-collapse:collapse">'
        f'<thead><tr style="background:#e0e0e0">{head}</tr></thead><tbody>{"".join(body)}</tbody></table>'
        f'{foot}<p style="color:#888; font-size:7pt">Gerado pelo IGNF Finora em '
        f'{generated:%d/%m/%Y %H:%M}</p></body></html>'
    )
